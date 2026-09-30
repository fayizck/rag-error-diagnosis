"""Read-only reconciliation of the two supplied human-review workbooks."""
from __future__ import annotations

import hashlib
import json
import sqlite3
from collections import Counter
from itertools import combinations
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[2]
OUT = Path(__file__).resolve().parent


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def sheet_rows(path):
    sheet = load_workbook(path, data_only=True)["Human Review"]
    names = [cell.value for cell in sheet[1]]
    rows = [dict(zip(names, row)) for row in sheet.iter_rows(min_row=2, values_only=True) if row[0]]
    assert len({r["Item ID"] for r in rows}) == len(rows)
    return rows


def observations(path):
    con = sqlite3.connect(path)
    try:
        return {sid: json.loads(payload) for sid, payload in con.execute("SELECT scientific_id,payload FROM observations")}
    finally:
        con.close()


def main():
    p1 = OUT / "evidence/study1_review.xlsx"
    p2 = OUT / "evidence/study2_review.xlsx"
    study1, study2 = sheet_rows(p1), sheet_rows(p2)
    assert len(study1) == 64 and len(study2) == 48
    retrieval = {r["item_id"]: r for r in map(json.loads, (ROOT / "outputs/study2_rag/retrieval_manifest.jsonl").read_text().splitlines())}
    prompts = {r["item_id"]: r for r in map(json.loads, (ROOT / "outputs/study2_rag/prompt_manifest.jsonl").read_text().splitlines())}
    s1 = observations(ROOT / "outputs/main/collection.sqlite3")
    s2 = observations(ROOT / "outputs/study2_rag/main/collection.sqlite3")
    conditions = ("00", "10", "01", "11")
    for r in study1:
        item = r["Item ID"]
        assert r["Question"] == retrieval[item]["question"] and r["Gold answer"] == retrieval[item]["answer"]
        for condition in conditions:
            assert r[f"{condition} response"] == s1[f"{item}::{condition}"]["answer"]
            assert r[f"{condition} semantic status"] in {"correct", "partially_correct", "incorrect", "ambiguous", "cannot_determine"}
        for a, b in combinations(conditions, 2):
            if r[f"{a} response"] == r[f"{b} response"]:
                assert r[f"{a} semantic status"] == r[f"{b} semantic status"]
    expected = {item for item, r in retrieval.items() if r["retrieval_category"] == "complete" and s2[f"{item}::rag"]["em"] == 0}
    assert {r["Item ID"] for r in study2} == expected
    for r in study2:
        item = r["Item ID"]
        assert r["Question"] == retrieval[item]["question"] and r["Gold answer"] == retrieval[item]["answer"]
        assert r["RAG response"] == s2[f"{item}::rag"]["answer"]
        assert r["Exact retrieved prompt"] == prompts[item]["prompt"]
        assert r["Answer semantic status"] in {"correct", "partially_correct", "incorrect", "ambiguous", "cannot_determine"}
        assert r["Context sufficiency"] in {"sufficient", "insufficient", "ambiguous", "cannot_determine"}
        assert r["Question/gold validity"] in {"valid", "questionable", "invalid", "cannot_determine"}
    first = {
        "selected_questions": len(study1),
        "selection": "Post-hoc lexical-review sample, not a random sample of all 200 questions",
        "semantic_status_by_condition": {condition: dict(Counter(r[f"{condition} semantic status"] for r in study1)) for condition in conditions},
        "official_em_correct_by_condition": {condition: sum(s1[f'{r["Item ID"]}::{condition}']["em"] for r in study1) for condition in conditions},
        "condition_00_to_11_status_transitions": {f"{a} -> {b}": n for (a, b), n in Counter((r["00 semantic status"], r["11 semantic status"]) for r in study1).items()},
        "condition_00_to_11_response_changes": sum(r["00 response"] != r["11 response"] for r in study1),
    }
    second = {
        "audited_cases": len(study2),
        "selection": "All official-EM failures among 106 questions with both annotated support titles in Study 2 top ten",
        "semantic_status": dict(Counter(r["Answer semantic status"] for r in study2)),
        "context_sufficiency": dict(Counter(r["Context sufficiency"] for r in study2)),
        "question_gold_validity": dict(Counter(r["Question/gold validity"] for r in study2)),
        "status_by_sufficiency": {f"{a} | {b}": n for (a, b), n in Counter((r["Answer semantic status"], r["Context sufficiency"]) for r in study2).items()},
        "incorrect_sufficient_review_labels": [r["Item ID"] for r in study2 if r["Answer semantic status"] == "incorrect" and r["Context sufficiency"] == "sufficient"],
        "incorrect_sufficient_valid_review_labels": [r["Item ID"] for r in study2 if r["Answer semantic status"] == "incorrect" and r["Context sufficiency"] == "sufficient" and r["Question/gold validity"] == "valid"],
        "reviewer_initials": sorted(set(r["Reviewer initials"] for r in study2)),
        "rows_with_notes": sum(bool(r["Human note"]) for r in study2),
    }
    result = {
        "study1_review_sha256": sha(p1), "study2_review_sha256": sha(p2),
        "study1": first, "study2": second,
        "qa_flags": {
            "5a84bc5b5542991dd0999dc3": "Question asks for a vizier after the 17-year vizier, but the gold Ibrium is described in the supplied text as his predecessor.",
            "5a8c5be0554299653c1aa04b": "Question omits 'in Europe'; supplied text calls Pedro Calomino the inventor and Law Adam the first European user.",
            "5a808f96554299485f598655": "Question asks for an event; 'First Balkan War' may be a defensible answer although the gold specifies women accompanying husbands in that war.",
        },
        "qa_flag_policy": "These are independent source-text checks, not silent edits to the human workbook. Treat the three flags as sensitivity exclusions for a strict unambiguous-failure count unless adjudicated.",
    }
    result["study2"]["strict_unambiguous_candidates_after_qa_flags"] = [
        item for item in second["incorrect_sufficient_valid_review_labels"] if item not in result["qa_flags"]
    ]
    (OUT / "reconciliation.json").write_text(json.dumps(result, ensure_ascii=False, indent=2, sort_keys=True) + "\n")
    print(json.dumps({"status": "PASS", "study1_rows": len(study1), "study2_rows": len(study2),
                      "study2_semantic_status": second["semantic_status"],
                      "study2_incorrect_sufficient": len(second["incorrect_sufficient_review_labels"]),
                      "qa_flags": len(result["qa_flags"])}, indent=2))


if __name__ == "__main__":
    main()
